"""在庫報告 PPT 生成。

画面 KPI カードを 1 カード 1 スライドで出力する。
表紙に「生産管理部」「在庫報告」を明示する。
"""
from __future__ import annotations

import base64
import io
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Sequence, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# 画面（月受注／在庫報告）に近いインディゴ系
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
INDIGO_DEEP = RGBColor(0x37, 0x30, 0xA3)
VIOLET = RGBColor(0x7C, 0x3A, 0xED)
VIOLET_SOFT = RGBColor(0xEE, 0xF2, 0xFF)
TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_SOFT = RGBColor(0xF0, 0xFD, 0xFA)
AMBER = RGBColor(0xD9, 0x77, 0x06)
AMBER_SOFT = RGBColor(0xFF, 0xFB, 0xEB)
ROSE = RGBColor(0xE1, 0x1D, 0x48)
ROSE_SOFT = RGBColor(0xFF, 0xF1, 0xF2)
EMERALD = RGBColor(0x05, 0x96, 0x69)
EMERALD_SOFT = RGBColor(0xEC, 0xFD, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_DARK = RGBColor(0x0F, 0x17, 0x2A)
TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)
SLATE = RGBColor(0x47, 0x55, 0x69)
CARD_LINE = RGBColor(0xE0, 0xE7, 0xFF)
FONT = "游ゴシック"

TARGET_SCRAP_RATE = 1.5
TARGET_PRODUCT_SEN = 375
TARGET_MATCH_RATE = 95.0


def _fmt_sen(v: Any) -> str:
    try:
        return f"{float(v) / 1000:,.1f}"
    except (TypeError, ValueError):
        return "—"


def _fmt_pct(v: Any, digits: int = 2) -> str:
    try:
        return f"{float(v):.{digits}f}%"
    except (TypeError, ValueError):
        return "—"


def _fmt_int(v: Any) -> str:
    try:
        return f"{int(round(float(v))):,}"
    except (TypeError, ValueError):
        return "—"


def _decode_chart_image(data_url: Optional[str]) -> Optional[bytes]:
    if not data_url:
        return None
    raw = str(data_url).strip()
    m = re.match(r"^data:image/(?:png|jpeg|jpg);base64,(.+)$", raw, re.I | re.S)
    payload = m.group(1) if m else raw
    try:
        return base64.b64decode(payload)
    except Exception:
        return None


def _new_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank_slide(prs: Presentation):
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    return prs.slides.add_slide(layout)


def _apply_ea_font(run, font_name: str = FONT) -> None:
    """日本語テキストが游ゴシックになるよう East Asian フォントも設定する。"""
    r_pr = run._r.get_or_add_rPr()
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    for local in ("latin", "ea", "cs"):
        tag = qn(f"a:{local}")
        el = r_pr.find(tag)
        if el is None:
            el = parse_xml(f'<{local} xmlns="{ns}" typeface="{font_name}"/>')
            r_pr.append(el)
        else:
            el.set("typeface", font_name)


def _set_run(run, *, size: int, bold: bool = False, color: RGBColor = TEXT_DARK) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _apply_ea_font(run, FONT)


def _add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    font_size: int = 14,
    bold: bool = False,
    color: RGBColor = TEXT_DARK,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=font_size, bold=bold, color=color)
    return box


def _add_rect(slide, left, top, width, height, fill: RGBColor, *, line: Optional[RGBColor] = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def _add_rounded(slide, left, top, width, height, fill: RGBColor, *, line: Optional[RGBColor] = None, adj: float = 0.15):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = adj
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def _fill_shape_text(
    shape,
    lines: Sequence[Tuple[str, int, bool, RGBColor]],
    *,
    align=PP_ALIGN.LEFT,
) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    if not lines:
        tf.text = ""
        return
    tf.text = lines[0][0]
    p0 = tf.paragraphs[0]
    p0.alignment = align
    p0.space_before = Pt(0)
    p0.space_after = Pt(1)
    for run in p0.runs:
        _set_run(run, size=lines[0][1], bold=lines[0][2], color=lines[0][3])
    for text, size, bold, color in lines[1:]:
        p = tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        run = p.add_run()
        run.text = text
        _set_run(run, size=size, bold=bold, color=color)


def _add_picture_fit(slide, png: bytes, left, top, max_w, max_h):
    """枠いっぱいに拡大（上下の余白を抑え、グラフを大きく表示）。"""
    stream = io.BytesIO(png)
    try:
        slide.shapes.add_picture(stream, left, top, width=max_w, height=max_h)
    except Exception:
        stream.seek(0)
        slide.shapes.add_picture(stream, left, top, width=max_w)


def _dept_badge(slide, left, top) -> None:
    badge = _add_rounded(slide, left, top, Inches(1.75), Inches(0.38), INDIGO, adj=0.5)
    _fill_shape_text(badge, [("生産管理部", 13, True, WHITE)], align=PP_ALIGN.CENTER)


def _build_cover(prs: Presentation, meta: Dict[str, Any]) -> None:
    slide = _blank_slide(prs)

    # 背景グラデ風（重ね矩形）
    _add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, INDIGO_DEEP)
    _add_rect(slide, Inches(0), Inches(0), Inches(8.2), SLIDE_H, INDIGO)
    accent = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.2), Inches(-1.2), Inches(5.5), Inches(5.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = VIOLET
    accent.line.fill.background()
    accent2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.6), Inches(4.2), Inches(4.2), Inches(4.2))
    accent2.fill.solid()
    accent2.fill.fore_color.rgb = RGBColor(0x63, 0x66, 0xF1)
    accent2.line.fill.background()

    # 白い内容カード
    card = _add_rounded(
        slide,
        Inches(0.85),
        Inches(1.35),
        Inches(8.4),
        Inches(4.7),
        WHITE,
        adj=0.06,
    )
    card.shadow.inherit = False

    _add_textbox(slide, Inches(1.2), Inches(1.7), Inches(7.6), Inches(0.35), "Smart-EMAPs", font_size=16, bold=True, color=INDIGO)
    pill = _add_rounded(slide, Inches(1.2), Inches(2.15), Inches(1.9), Inches(0.4), INDIGO, adj=0.5)
    _fill_shape_text(pill, [("生産管理部", 14, True, WHITE)], align=PP_ALIGN.CENTER)

    _add_textbox(slide, Inches(1.2), Inches(2.7), Inches(7.6), Inches(0.65), "在庫報告", font_size=44, bold=True, color=TEXT_DARK)
    _add_textbox(
        slide,
        Inches(1.2),
        Inches(3.35),
        Inches(7.6),
        Inches(0.4),
        str(meta.get("label") or "在庫報告書"),
        font_size=24,
        bold=True,
        color=INDIGO,
    )
    period = f"対象期間：{meta.get('period_start') or '—'} 〜 {meta.get('period_end') or '—'}"
    _add_textbox(slide, Inches(1.2), Inches(3.9), Inches(7.6), Inches(0.35), period, font_size=15, color=TEXT_MUTED)

    # 主要指標チップ
    chips = [
        ("月末在庫", INDIGO),
        ("廃棄率", AMBER),
        ("廃棄本数", ROSE),
        ("棚卸一致", EMERALD),
        ("工程差異", VIOLET),
        ("大量廃棄", TEAL),
    ]
    chip_w = 1.15
    chip_gap = 0.08
    chip_start = 1.2
    for i, (label, color) in enumerate(chips):
        left = Inches(chip_start + i * (chip_w + chip_gap))
        chip = _add_rounded(slide, left, Inches(4.55), Inches(chip_w), Inches(0.42), color, adj=0.45)
        _fill_shape_text(chip, [(label, 12, True, WHITE)], align=PP_ALIGN.CENTER)

    # 右側キャッチ
    _add_textbox(
        slide,
        Inches(9.5),
        Inches(2.6),
        Inches(3.4),
        Inches(0.4),
        "PRODUCTION",
        font_size=14,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.LEFT,
    )
    _add_textbox(
        slide,
        Inches(9.5),
        Inches(3.05),
        Inches(3.4),
        Inches(0.9),
        "Inventory\nReport",
        font_size=30,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.LEFT,
    )
    _add_textbox(
        slide,
        Inches(9.5),
        Inches(4.3),
        Inches(3.4),
        Inches(0.8),
        "月末在庫・廃棄率・棚卸差異\nを報告向けに集約",
        font_size=14,
        color=RGBColor(0xE0, 0xE7, 0xFF),
    )


def _content_chrome(
    slide,
    *,
    title: str,
    subtitle: str = "",
    accent: RGBColor = INDIGO,
    page_no: int = 1,
    page_total: int = 1,
    meta_label: str = "",
) -> None:
    # 背景
    _add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, RGBColor(0xF8, 0xFA, 0xFC))
    # 上部バー
    _add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.72), accent)
    _dept_badge(slide, Inches(0.35), Inches(0.16))
    _add_textbox(
        slide,
        Inches(2.25),
        Inches(0.16),
        Inches(7.3),
        Inches(0.4),
        "在庫報告",
        font_size=16,
        bold=True,
        color=WHITE,
    )
    if meta_label:
        _add_textbox(
            slide,
            Inches(8.6),
            Inches(0.18),
            Inches(4.3),
            Inches(0.38),
            meta_label,
            font_size=13,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.RIGHT,
        )
    # 左アクセント
    _add_rect(slide, Inches(0), Inches(0.72), Inches(0.1), Inches(6.78), accent)
    # タイトル
    _add_textbox(slide, Inches(0.4), Inches(0.88), Inches(12.4), Inches(0.48), title, font_size=26, bold=True, color=TEXT_DARK)
    if subtitle:
        _add_textbox(slide, Inches(0.4), Inches(1.32), Inches(12.4), Inches(0.32), subtitle, font_size=14, color=TEXT_MUTED)
    # フッター
    _add_rect(slide, Inches(0), Inches(7.15), SLIDE_W, Inches(0.35), WHITE)
    _add_textbox(
        slide,
        Inches(0.4),
        Inches(7.17),
        Inches(8),
        Inches(0.3),
        "生産管理部 ／ Smart-EMAPs 在庫報告管理",
        font_size=11,
        color=TEXT_MUTED,
    )
    _add_textbox(
        slide,
        Inches(10.5),
        Inches(7.17),
        Inches(2.4),
        Inches(0.3),
        f"{page_no} / {page_total}",
        font_size=11,
        bold=True,
        color=TEXT_MUTED,
        align=PP_ALIGN.RIGHT,
    )


def _month_chips(
    slide,
    months: Sequence[Dict[str, Any]],
    value_fn,
    *,
    top: float = 1.65,
    accent: RGBColor = INDIGO,
) -> float:
    """月次 KPI チップを描画し、チップ下端 Y（インチ）を返す。"""
    if not months:
        return top
    n = len(months)
    gap = 0.08 if n >= 10 else 0.1
    usable = 12.4
    chip_w = min(1.65, (usable - gap * max(n - 1, 0)) / max(n, 1))
    total = n * chip_w + (n - 1) * gap
    start = 0.45 + max(0, (usable - total) / 2)
    label_sz = 9 if n >= 10 else 10
    value_sz = 12 if n >= 10 else 14
    chip_h = 0.68
    for i, m in enumerate(months):
        left = Inches(start + i * (chip_w + gap))
        bar = _add_rounded(slide, left, Inches(top), Inches(chip_w), Inches(0.06), accent, adj=0.5)
        del bar
        card = _add_rounded(
            slide,
            left,
            Inches(top + 0.04),
            Inches(chip_w),
            Inches(chip_h),
            WHITE,
            line=CARD_LINE,
            adj=0.18,
        )
        label = str(m.get("month_label") or m.get("month") or "")
        value = value_fn(m)
        _fill_shape_text(
            card,
            [(label, label_sz, True, TEXT_MUTED), (value, value_sz, True, accent)],
            align=PP_ALIGN.CENTER,
        )
    return top + 0.04 + chip_h


def _kpi_card_slide(
    prs: Presentation,
    *,
    title: str,
    subtitle: str,
    months: Sequence[Dict[str, Any]],
    value_fn,
    chart_png: Optional[bytes],
    accent: RGBColor,
    page_no: int,
    page_total: int,
    meta_label: str,
    show_months: bool = True,
) -> None:
    slide = _blank_slide(prs)
    _content_chrome(
        slide,
        title=title,
        subtitle=subtitle,
        accent=accent,
        page_no=page_no,
        page_total=page_total,
        meta_label=meta_label,
    )
    if show_months:
        chips_bottom = _month_chips(slide, months, value_fn, top=1.65, accent=accent)
        chart_top = chips_bottom + 0.12
    else:
        chart_top = 1.68

    chart_panel = _add_rounded(
        slide,
        Inches(0.4),
        Inches(chart_top),
        Inches(12.55),
        Inches(6.95 - chart_top),
        WHITE,
        line=CARD_LINE,
        adj=0.05,
    )
    del chart_panel
    if chart_png:
        _add_picture_fit(
            slide,
            chart_png,
            Inches(0.48),
            Inches(chart_top + 0.06),
            Inches(12.4),
            Inches(6.95 - chart_top - 0.1),
        )
    else:
        _add_textbox(
            slide,
            Inches(0.65),
            Inches(4.4),
            Inches(12.05),
            Inches(0.4),
            "グラフデータがありません",
            font_size=16,
            color=TEXT_MUTED,
            align=PP_ALIGN.CENTER,
        )


def _diff_slide(
    prs: Presentation,
    *,
    chart_png: Optional[bytes],
    page_no: int,
    page_total: int,
    meta_label: str,
    title: str = "理論在庫 vs 棚卸差異（工程別）",
    subtitle: str = "対象期間の月平均・単位：千本（検査は倉庫へ合算）",
    accent: RGBColor = VIOLET,
) -> None:
    """差異スライドはグラフのみ（下部データ表なし）。"""
    slide = _blank_slide(prs)
    _content_chrome(
        slide,
        title=title,
        subtitle=subtitle,
        accent=accent,
        page_no=page_no,
        page_total=page_total,
        meta_label=meta_label,
    )
    chart_top = 1.68
    panel = _add_rounded(
        slide,
        Inches(0.4),
        Inches(chart_top),
        Inches(12.55),
        Inches(6.95 - chart_top),
        WHITE,
        line=CARD_LINE,
        adj=0.05,
    )
    del panel
    if chart_png:
        _add_picture_fit(
            slide,
            chart_png,
            Inches(0.48),
            Inches(chart_top + 0.06),
            Inches(12.4),
            Inches(6.95 - chart_top - 0.1),
        )
    else:
        _add_textbox(
            slide,
            Inches(0.65),
            Inches(4.2),
            Inches(12.05),
            Inches(0.4),
            "グラフデータがありません",
            font_size=16,
            color=TEXT_MUTED,
            align=PP_ALIGN.CENTER,
        )


def _bulk_chart_slide(
    prs: Presentation,
    *,
    title: str,
    subtitle: str,
    chart_png: Optional[bytes],
    categories: Sequence[Dict[str, Any]],
    page_no: int,
    page_total: int,
    meta_label: str,
    show_cats: bool = True,
) -> None:
    slide = _blank_slide(prs)
    _content_chrome(
        slide,
        title=title,
        subtitle=subtitle,
        accent=AMBER,
        page_no=page_no,
        page_total=page_total,
        meta_label=meta_label,
    )
    y_chart = 1.7
    if show_cats and categories:
        chips_bottom = _month_chips(
            slide,
            [
                {
                    "month_label": str(c.get("name") or ""),
                    "_qty": c.get("quantity"),
                }
                for c in categories[:8]
            ],
            lambda m: _fmt_int(m.get("_qty")),
            top=1.65,
            accent=AMBER,
        )
        y_chart = chips_bottom + 0.12

    panel = _add_rounded(
        slide,
        Inches(0.4),
        Inches(y_chart),
        Inches(12.55),
        Inches(6.9 - y_chart - 0.35),
        WHITE,
        line=CARD_LINE,
        adj=0.05,
    )
    del panel
    if chart_png:
        _add_picture_fit(
            slide,
            chart_png,
            Inches(0.65),
            Inches(y_chart + 0.15),
            Inches(12.05),
            Inches(6.9 - y_chart - 0.6),
        )
    else:
        _add_textbox(
            slide,
            Inches(0.65),
            Inches(4.2),
            Inches(12.05),
            Inches(0.4),
            "グラフデータがありません",
            font_size=16,
            color=TEXT_MUTED,
            align=PP_ALIGN.CENTER,
        )


def build_inventory_report_pptx(body: Dict[str, Any]) -> bytes:
    """在庫報告 PPT（表紙 + KPIカード各1ページ）。"""
    payload = body.get("payload") or {}
    months = body.get("monthly_kpis") or payload.get("monthly_kpis") or []
    charts_raw = body.get("charts") or {}

    charts = {k: _decode_chart_image(v) for k, v in charts_raw.items()}

    label = str(payload.get("label") or "在庫報告")
    period = payload.get("period") or {}
    meta = {
        "label": label,
        "period_start": period.get("start"),
        "period_end": period.get("end"),
        "generated_at": payload.get("generated_at") or "—",
        "exported_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

    # ページ構成を先に数え上げ
    content_specs: List[str] = [
        "inventory",
        "wip",
        "product",
        "scrap_rate",
        "scrap_qty",
        "match",
        "diff",
    ]
    page_total = 1 + len(content_specs)  # cover + contents
    prs = _new_prs()
    _build_cover(prs, meta)

    page = 2
    _kpi_card_slide(
        prs,
        title="月末在庫（仕掛品・製品）",
        subtitle="仕掛品＋製品の月末在庫推移",
        months=months,
        value_fn=lambda m: _fmt_sen(m.get("closing_total_qty")),
        chart_png=charts.get("inventory"),
        accent=INDIGO,
        page_no=page,
        page_total=page_total,
        meta_label=label,
        show_months=False,
    )
    page += 1

    _kpi_card_slide(
        prs,
        title="期末仕掛品",
        subtitle="",
        months=months,
        value_fn=lambda m: _fmt_sen(m.get("closing_wip_qty")),
        chart_png=charts.get("wip"),
        accent=RGBColor(0x25, 0x63, 0xEB),
        page_no=page,
        page_total=page_total,
        meta_label=label,
    )
    page += 1

    _kpi_card_slide(
        prs,
        title=f"期末製品　目安 {TARGET_PRODUCT_SEN}千本",
        subtitle="グラフ破線＝管理目安",
        months=months,
        value_fn=lambda m: _fmt_sen(m.get("closing_product_qty")),
        chart_png=charts.get("product"),
        accent=TEAL,
        page_no=page,
        page_total=page_total,
        meta_label=label,
    )
    page += 1

    _kpi_card_slide(
        prs,
        title=f"廃棄率　目安 {TARGET_SCRAP_RATE}%",
        subtitle="主ライン（切断〜検査）　グラフ破線＝管理目安",
        months=months,
        value_fn=lambda m: _fmt_pct(m.get("quality_loss_rate_percent")),
        chart_png=charts.get("scrapRate") or charts.get("scrap_rate"),
        accent=AMBER,
        page_no=page,
        page_total=page_total,
        meta_label=label,
    )
    page += 1

    _kpi_card_slide(
        prs,
        title="廃棄本数",
        subtitle="不良＋廃棄の合計",
        months=months,
        value_fn=lambda m: _fmt_sen(
            m.get("loss_qty")
            if m.get("loss_qty") is not None
            else (float(m.get("defect_qty") or 0) + float(m.get("scrap_qty") or 0))
        ),
        chart_png=charts.get("scrapQty") or charts.get("scrap_qty"),
        accent=ROSE,
        page_no=page,
        page_total=page_total,
        meta_label=label,
    )
    page += 1

    _kpi_card_slide(
        prs,
        title=f"棚卸一致率　目安 {TARGET_MATCH_RATE:.0f}%",
        subtitle="製品×工程の一致件数割合　グラフ破線＝管理目安",
        months=months,
        value_fn=lambda m: _fmt_pct(m.get("match_rate")),
        chart_png=charts.get("match"),
        accent=EMERALD,
        page_no=page,
        page_total=page_total,
        meta_label=label,
    )
    page += 1

    _diff_slide(
        prs,
        chart_png=charts.get("diff"),
        page_no=page,
        page_total=page_total,
        meta_label=label,
    )

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
